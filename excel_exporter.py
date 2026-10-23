"""Seven-sheet industry-grouped export matching the user's sample layout.

Missing data remains blank. Never infer that no recorded event means zero events.
"""
import csv
import calendar
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from paths import DATA_DIR, OUTPUT_DIR, ROOT
from stock_data import number
from split_history import saved_split_events

SHEETS = ('Stock Selection', 'Price Data', 'Cash Dividend', 'Stock Dividend',
          'Right Share Ratio', 'Right Share Offer Price', 'Split')
DARK = '21334D'
BLUE = 'DCEAF6'
GRAY = 'F1F5FA'
DIVIDEND_COLOR = 'F5FAFF'


def rows_from_csv(path):
    if not path.exists():
        return []
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def read_corporate_actions(path=None):
    """Read user-confirmed corporate actions. Ratio means new/old shares for rights/splits;
    stock dividend is a percentage of existing shares (10 = 10 percent).
    """
    path = Path(path) if path else ROOT / 'corporate_actions.csv'
    events = []
    for lineno, row in enumerate(rows_from_csv(path), 2):
        ticker = (row.get('ticker') or '').strip().upper()
        kind = (row.get('event_type') or '').strip().lower()
        raw_date = (row.get('date') or '').strip()
        if not ticker or not kind or not raw_date:
            continue
        if kind not in ('stock_dividend', 'rights', 'split'):
            raise ValueError(f'{path.name}:{lineno}: invalid event_type: {kind!r}')
        try:
            dt = date.fromisoformat(raw_date)
        except ValueError as exc:
            raise ValueError(f'{path.name}:{lineno}: date must be YYYY-MM-DD') from exc
        raw_ratio = (row.get('ratio') or '').strip()
        ratio = parse_ratio(raw_ratio, kind) if raw_ratio else None
        offer = number(row.get('offer_price'))
        if kind in ('stock_dividend', 'split') and (ratio is None or ratio <= 0):
            raise ValueError(f'{path.name}:{lineno}: {kind} requires positive ratio')
        if kind == 'rights' and (ratio is None or ratio <= 0):
            raise ValueError(f'{path.name}:{lineno}: rights requires positive ratio')
        if kind == 'rights' and (offer is None or offer < 0):
            # Ratio may be known while offer price is not yet known.
            offer = None
        events.append({'ticker': ticker, 'kind': kind, 'date': dt, 'ratio': ratio,
                       'offer_price': offer, 'source_url': (row.get('source_url') or '').strip()})
    return events


def parse_ratio(raw, kind):
    """1R:2 -> 0.5 rights/share; 2:1 split -> 2; stock 10% -> 10 percentage points."""
    text = raw.strip().upper().replace(' ', '')
    if ':' in text:
        bits = text.replace('R','').split(':')
        if len(bits) != 2:
            return None
        a, b = number(bits[0]), number(bits[1])
        return a / b if a is not None and b not in (None, 0) else None
    val = number(text)
    if kind == 'stock_dividend' and text.endswith('%') and val is not None:
        return val * 100
    return val


def by_month(rows, date_field, field, start, end, aggregate='last'):
    """Retain last available close in each month; sum multiple cash dividends."""
    output = {}
    ordered = sorted(rows, key=lambda r: (r.get(date_field) or ''))
    for r in ordered:
        try:
            dt = date.fromisoformat(str(r[date_field])[:10])
        except (ValueError, KeyError, TypeError):
            continue
        if dt < start or dt > end:
            continue
        val = number(r.get(field))
        if val is None:
            continue
        k = (dt.year, dt.month)
        if aggregate == 'sum':
            output[k] = output.get(k, 0) + val
        elif aggregate == 'product':
            output[k] = output.get(k, 1) * val
        else:
            output[k] = val
    return output


def _company_info(tickers):
    infos = {}
    for ticker in tickers:
        path = DATA_DIR / ticker / 'company.json'
        try:
            info = json.loads(path.read_text(encoding='utf-8'))
        except (OSError, ValueError):
            info = {}
        infos[ticker] = {'ticker': ticker, 'name': str(info.get('name') or ticker),
                          'industry': str(info.get('industry') or 'Unknown Industry')}
    return infos


def _sort_tickers(tickers, info):
    return sorted(tickers, key=lambda t: (info[t]['industry'] == 'Unknown Industry',
                                        info[t]['industry'].casefold(), info[t]['name'].casefold(), t))


def _apply_sheet_basics(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.freeze_panes = 'C3'
    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['B'].width = 15
    ws.row_dimensions[1].height = 27
    ws.row_dimensions[2].height = 29


def _write_monthly(ws, tickers, info, start, end, values, number_format='0.00', source_note=''):
    _apply_sheet_basics(ws)
    ws['A2'] = 'Years'
    ws['B2'] = 'Months'
    header_fill = PatternFill('solid', fgColor=DARK)
    for col in ('A2', 'B2'):
        cell=ws[col]
        cell.fill = header_fill
        cell.font = Font(name='Aptos', bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    current_industry = None
    group_start = 3
    for index, ticker in enumerate(tickers):
        col = index + 3
        industry = info[ticker]['industry']
        letter = get_column_letter(col)
        ws.cell(2, col, info[ticker]['name'])
        h=ws.cell(2, col)
        h.font=Font(name='Aptos', size=10, bold=True, color='FFFFFF')
        h.fill=header_fill
        h.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
        h.comment = Comment(f"Ticker: {ticker}\nIndustry: {industry}\n{source_note}", 'DSE Scraper')
        ws.column_dimensions[letter].width = max(16, min(28, len(info[ticker]['name'])+4))
        if current_industry is None:
            current_industry = industry
            group_start=col
        elif industry != current_industry:
            _finish_industry(ws, current_industry, group_start, col-1)
            current_industry, group_start = industry, col
    if current_industry is not None:
        _finish_industry(ws, current_industry, group_start, len(tickers)+2)

    row=3
    last_year = None
    for year in range(start.year, end.year+1):
        for month in range(1, 13):
            if (year, month) < (start.year, start.month) or (year, month) > (end.year, end.month):
                continue
            ws.cell(row,1,year if last_year != year else None)
            last_year = year
            ws.cell(row,2,calendar.month_name[month])
            for idx,ticker in enumerate(tickers,3):
                value=values.get(ticker,{}).get((year,month))
                if value is not None:
                    c=ws.cell(row,idx,value)
                    c.number_format=number_format
            fill=PatternFill('solid',fgColor='FFFFFF' if row%2 else GRAY)
            for col in range(1,len(tickers)+3):
                c=ws.cell(row,col)
                c.fill=fill
                c.font=Font(name='Aptos',size=10,color=DARK)
                c.alignment=Alignment(vertical='center',horizontal='left' if col==2 else 'right')
            if ws.cell(row,1).value is not None:
                ws.cell(row,1).font=Font(name='Aptos',bold=True,color=DARK)
            ws.row_dimensions[row].height=19
            row+=1
    ws.auto_filter.ref = f'B2:{get_column_letter(len(tickers)+2)}{row-1}' if tickers else None
    ws.print_title_rows = '1:2'


def _finish_industry(ws, name, first, last):
    if first < last:
        ws.merge_cells(start_row=1,start_column=first,end_row=1,end_column=last)
    c=ws.cell(1,first,name)
    c.fill=PatternFill('solid',fgColor=BLUE)
    c.font=Font(name='Aptos',bold=True,color=DARK)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for col in range(first,last+1):
        ws.cell(1,col).fill=PatternFill('solid',fgColor=BLUE)


def _write_ratios(ws, tickers, info, start, end):
    ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=19
    for letter in 'BCDEFG':ws.column_dimensions[letter].width=15
    ws.freeze_panes='B3'
    row=1
    for index,ticker in enumerate(tickers,1):
        title = ws.cell(row,1,f'{index}. {info[ticker]["name"]}')
        title.font=Font(name='Aptos',size=12,bold=True,color='B92336')
        title.comment=Comment(f'Ticker: {ticker}\nFinancials: https://stockanalysis.com/quote/dse/{ticker}/financials/ratios/\nBook value/share: https://stockanalysis.com/quote/dse/{ticker}/financials/balance-sheet/', 'DSE Scraper')
        ws.row_dimensions[row].height=24
        row+=1
        for col,head in enumerate(('Year','P/E','NAV','ROA','ROE','NPM','EPS'),1):
            c=ws.cell(row,col,head)
            c.fill=PatternFill('solid',fgColor=DARK)
            c.font=Font(name='Aptos',size=10,color='FFFFFF',bold=True)
            c.alignment=Alignment(horizontal='center')
        row+=1
        source=DATA_DIR/ticker/f'{ticker}_ratios_export.csv'
        years={int(d['year']):d for d in rows_from_csv(source) if str(d.get('year','')).isdigit()}
        for year in range(start.year, end.year+1):
            ws.cell(row,1,year)
            source_row=years.get(year,{})
            for col,key in enumerate(('pe','nav','roa','roe','npm','eps'),2):
                value=number(source_row.get(key))
                if value is not None:
                    c=ws.cell(row,col,value)
                    c.number_format='0.00%' if key in ('roa','roe','npm') else '0.00'
            for col in range(1,8):
                c=ws.cell(row,col)
                c.fill=PatternFill('solid',fgColor='FFFFFF' if row%2 else GRAY)
                c.font=Font(name='Aptos',size=10,color=DARK)
            row+=1
        row+=1


def create_excel(start_date=None, end_date=None, tickers=None, output_path=None, actions_path=None):
    if isinstance(start_date,int):start_date=date(start_date,1,1)
    if isinstance(end_date,int):end_date=date(end_date,12,31)
    start_date = start_date or date(date.today().year-4,1,1)
    end_date = end_date or date.today()
    if isinstance(start_date,datetime):start_date=start_date.date()
    if isinstance(end_date,datetime):end_date=end_date.date()
    if start_date>end_date:raise ValueError('Start date must be before end date')
    if tickers is None:
        tickers=sorted(p.name for p in DATA_DIR.iterdir() if p.is_dir()) if DATA_DIR.exists() else []
    tickers=list(dict.fromkeys(tickers))
    if not tickers:raise ValueError('No tickers supplied')
    info=_company_info(tickers)
    ordered=_sort_tickers(tickers,info)
    actions=read_corporate_actions(actions_path)
    # User-supplied/verified CSV takes precedence for the same ticker/event/date.
    # Do not multiply matching imported and scraped split ratios twice.
    overrides = {(e['ticker'], e['kind'], e['date']) for e in actions}
    actions.extend(e for e in saved_split_events(ordered)
                   if (e['ticker'], e['kind'], e['date']) not in overrides)

    price, cash, stock, right, offer, split = [defaultdict(dict) for _ in range(6)]
    for ticker in ordered:
        base=DATA_DIR/ticker
        price[ticker]=by_month(rows_from_csv(base/f'{ticker}_history.csv'),'date','close',start_date,end_date)
        cash[ticker]=by_month(rows_from_csv(base/f'{ticker}_dividend.csv'),'date','cash',start_date,end_date,'sum')
    provenance = defaultdict(list)
    for event in actions:
        ticker=event['ticker']
        dt=event['date']
        if ticker not in info or not start_date<=dt<=end_date:
            continue
        key=(dt.year,dt.month)
        kind=event['kind']
        val=event['ratio']
        detail = f"{dt.isoformat()} {kind}: ratio={val}; offer price={event['offer_price']}; source={event['source_url'] or 'not supplied'}"
        if kind == 'stock_dividend': provenance[(SHEETS[3], ticker, key)].append(detail)
        if kind == 'rights':
            provenance[(SHEETS[4], ticker, key)].append(detail)
            if event['offer_price'] is not None: provenance[(SHEETS[5], ticker, key)].append(detail)
        if kind == 'split': provenance[(SHEETS[6], ticker, key)].append(detail)
        if kind=='stock_dividend':
            stock[ticker][key]=stock[ticker].get(key,0)+val / 100
        elif kind=='rights':
            right[ticker][key]=right[ticker].get(key,0)+val
            if event['offer_price'] is not None:
                offer[ticker][key]=event['offer_price']
        elif kind=='split':
            split[ticker][key]=split[ticker].get(key,1)*val

    book=Workbook()
    book.remove(book.active)
    for sheet in SHEETS:book.create_sheet(sheet)
    _write_ratios(book[SHEETS[0]],ordered,info,start_date,end_date)
    layouts=[(SHEETS[1],price,'0.00','Month-end closing price or last available trading close. StockAnalysis history.'),
             (SHEETS[2],cash,'0.00000','Cash dividends in BDT per share; ex-dividend month. StockAnalysis dividend history.'),
             (SHEETS[3],stock,'0.00%','Confirmed stock dividends from corporate_actions.csv; month of supplied date.'),
             (SHEETS[4],right,'0.0000','Confirmed rights new shares per existing share from corporate_actions.csv.'),
             (SHEETS[5],offer,'0.00','Confirmed rights offer price (BDT/share) from corporate_actions.csv.'),
             (SHEETS[6],split,'0.0000','Split multiplier (new shares per old share) from StockAnalysis latest and verified history; blank=unverified.')]
    for name,values,fmt,note in layouts:
        _write_monthly(book[name],ordered,info,start_date,end_date,values,fmt,note)
    # Original source URLs stay attached to imported action values for auditing.
    ticker_column = {ticker:idx+3 for idx,ticker in enumerate(ordered)}
    month_row = {}
    current_row = 3
    for year in range(start_date.year,end_date.year+1):
        for month in range(1,13):
            if (year,month) < (start_date.year,start_date.month) or (year,month) > (end_date.year,end_date.month):
                continue
            month_row[(year,month)]=current_row
            current_row+=1
    for (sheet,ticker,month),details in provenance.items():
        book[sheet].cell(month_row[month],ticker_column[ticker]).comment=Comment('\n'.join(details),'DSE Scraper')
    output_path=Path(output_path) if output_path else OUTPUT_DIR/'portfolio.xlsx'
    output_path.parent.mkdir(parents=True,exist_ok=True)
    book.save(output_path)
    return output_path


if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser()
    p.add_argument('--start-year',type=int,default=date.today().year-4)
    p.add_argument('--end-year',type=int,default=date.today().year)
    p.add_argument('--tickers',required=True,help='Comma separated trading codes (using existing cached data)')
    p.add_argument('--output',default=str(OUTPUT_DIR/'portfolio.xlsx'))
    args=p.parse_args()
    result=create_excel(args.start_year,args.end_year,[x.strip().upper() for x in args.tickers.split(',') if x.strip()],args.output)
    print('Saved',result)
