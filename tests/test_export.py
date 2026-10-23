import csv
import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
import excel_exporter as ex
from stock_data import financial_table, ratio_data, company_metadata, history_data, dividend_data
from crawler import clean_ticker, filtered_data


def csv_rows(path,fieldnames,rows):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.data=self.root/'data'
        self.actions=self.root/'corporate_actions.csv'
        self.patcher=patch.object(ex,'DATA_DIR',self.data)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        self.tickers=['ZZBANK','AAFOOD','CCBANK']
        for ticker,industry,name in [('ZZBANK','Commercial Banks','Zeta Bank'),
                                     ('AAFOOD','Food Products','Alpha Foods'),
                                     ('CCBANK','Commercial Banks','City Bank')]:
            folder=self.data/ticker
            folder.mkdir(parents=True)
            (folder/'company.json').write_text(json.dumps({'ticker':ticker,'industry':industry,'name':name}))
            csv_rows(folder/f'{ticker}_ratios_export.csv',['year','pe','nav','roa','roe','npm','eps'],[
                {'year':2024,'pe':10,'nav':23,'roa':0.015,'roe':0.20,'npm':0.11,'eps':4.5}])
            csv_rows(folder/f'{ticker}_history.csv',['date','close'],[
                {'date':'2024-01-05','close':20},
                {'date':'2024-01-31','close':25},
                {'date':'2024-02-01','close':30},
                {'date':'2025-01-01','close':40}])
            csv_rows(folder/f'{ticker}_dividend.csv',['date','cash'],[
                {'date':'2024-01-15','cash':0.50},
                {'date':'2024-01-19','cash':0.75}])
        csv_rows(self.actions,['ticker','event_type','date','ratio','offer_price','source_url','notes'],[
            {'ticker':'ZZBANK','event_type':'stock_dividend','date':'2024-01-06','ratio':'10%','offer_price':'','source_url':'','notes':''},
            {'ticker':'ZZBANK','event_type':'rights','date':'2024-01-08','ratio':'1R:2','offer_price':'20','source_url':'','notes':''},
            {'ticker':'ZZBANK','event_type':'split','date':'2024-02-09','ratio':'2:1','offer_price':'','source_url':'','notes':''},
        ])

    def test_seven_sheets_industry_headers_ratios_and_corporate_actions(self):
        target=self.root/'out.xlsx'
        ex.create_excel(2024,2025,self.tickers,target,actions_path=self.actions)
        wb=load_workbook(target)
        self.assertEqual(wb.sheetnames,list(ex.SHEETS))
        # Banks sorted together (City then Zeta), then Food.
        for sheet in list(ex.SHEETS)[1:]:
            ws=wb[sheet]
            self.assertEqual([ws.cell(2,c).value for c in (3,4,5)],['City Bank','Zeta Bank','Alpha Foods'])
            self.assertEqual([ws.cell(1,c).value for c in (3,5)],['Commercial Banks','Food Products'])
            self.assertEqual(ws.max_row,26)  # 2 heading + 24 months
        price=wb['Price Data']
        self.assertEqual(price['D3'].value,25)
        self.assertEqual(price['D4'].value,30)
        self.assertEqual(price['A3'].value,2024)
        self.assertEqual(price['A4'].value,None)
        self.assertEqual(price['A15'].value,2025)
        self.assertEqual(wb['Cash Dividend']['D3'].value,1.25)
        self.assertAlmostEqual(wb['Stock Dividend']['D3'].value,0.1)
        self.assertEqual(wb['Right Share Ratio']['D3'].value,0.5)
        self.assertEqual(wb['Right Share Offer Price']['D3'].value,20)
        self.assertEqual(wb['Split']['D4'].value,2)
        self.assertIsNone(wb['Split']['D3'].value)
        ratios=wb['Stock Selection']
        self.assertEqual(ratios['A1'].value,'1. City Bank')
        self.assertEqual(ratios['E3'].value,0.2)
        self.assertEqual(ratios['E3'].number_format,'0.00%')

    def test_date_filter_and_no_fabricated_events(self):
        out=ex.create_excel(date(2024,2,1),date(2024,3,15),self.tickers,self.root/'short.xlsx',self.actions)
        wb=load_workbook(out)
        self.assertEqual(wb['Price Data'].max_row,4)
        self.assertEqual(wb['Price Data']['A3'].value,2024)
        self.assertEqual(wb['Price Data']['B3'].value,'February')
        self.assertIsNone(wb['Cash Dividend']['D3'].value)
        self.assertIsNone(wb['Split']['D4'].value)

    def test_validate_bad_event(self):
        csv_rows(self.actions,['ticker','event_type','date','ratio','offer_price','source_url','notes'],[
            {'ticker':'ZZBANK','event_type':'rights','date':'2024-01-05','ratio':'0','offer_price':'20'}])
        with self.assertRaises(ValueError):ex.read_corporate_actions(self.actions)


class SourceParsingTests(unittest.TestCase):
    def test_ticker_validation(self):
        self.assertEqual(clean_ticker('  bracbank '),'BRACBANK')
        for ticker in ['../cookies.json','ABC/DEF','','F|BAR']:
            with self.assertRaises(ValueError):clean_ticker(ticker)

    def test_auth_payload_removed(self):
        payload={'nodes':[{'data':['SECRET ACCESS TOKEN']},{'data':[{'info':1},{'name':2},'Company']},{'data':[{'infoTable':1},[]]}]}
        result=filtered_data(payload)
        self.assertIsNone(result['nodes'][0])
        self.assertNotIn('SECRET',json.dumps(result))

    def test_fiscal_excludes_ttm_and_calculates_npm(self):
        def table(section,items):
            d=[None]*70
            d[0]={'financialData':8}
            d[8]={}
            index=10
            for key,arr in items.items():
                ptr=index
                index+=1
                d[ptr]=list(range(index,index+len(arr)))
                for pos,item in enumerate(arr):d[index+pos]=item
                index+=len(arr)
                d[8][key]=ptr
            return {'url':f'https://stockanalysis.com/quote/dse/TEST/{section}/__data.json',
                    'response_body':{'nodes':[None,{}, {'data':d}]}}
        ratios=table('financials/ratios',{'datekey':['TTM','2024-12-31'],'fiscalYear':[2025,2024],
                                        'pe':[7,8],'roa':[0.3,0.01],'roe':[0.4,0.14]})
        income=table('financials',{'datekey':['2024-12-31'],'fiscalYear':[2024],
                                    'revenue':[1000],'netinccmn':[120],'epsdil':[2.5]})
        balance=table('financials/balance-sheet',{'datekey':['2024-12-31'],'fiscalYear':[2024],
                                                  'bvps':[25]})
        result=ratio_data([ratios,income,balance])
        self.assertNotIn(2025,result)
        self.assertEqual(result[2024]['pe'],8)
        self.assertEqual(result[2024]['nav'],25)
        self.assertAlmostEqual(result[2024]['npm'],0.12)



class InvestingSplitTests(unittest.TestCase):
    def test_split_parser_reads_fractional_and_large_ratios(self):
        from split_history import parse_split_history_text
        page = '''BSRM Steels (BSRM) Stock Split History
        Split date       Split ratio
        Oct 22, 2018    1.1:1 Stock Split
        Oct 09, 2017    1.1:1 Stock Split
        Mar 13, 2013    1.05:1 Stock Split
        Dec 04, 2011    10:1 Stock Split
        Apr 13, 2011    1.2:1 Stock Split
        May 06, 2010    1.15:1 Stock Split'''
        result = parse_split_history_text(page)
        self.assertEqual(len(result), 6)
        self.assertEqual(result[-1], {'date':'2018-10-22','ratio':'1.1'})
        self.assertEqual(next(x['ratio'] for x in result if x['date']=='2011-12-04'), '10')
        self.assertEqual(parse_split_history_text('A page containing 2:1 but no split table'),[])

    def test_identity_check_rejects_other_companies(self):
        from split_history import heading_matches_company, candidate_urls, split_url_mapping
        self.assertTrue(heading_matches_company('BSRM Steels (BSRM) Stock Split History', 'BSRM Steels Limited'))
        self.assertFalse(heading_matches_company('BSRM Holdings (BSRM) Stock Split History', 'BSRM Steels Limited'))
        urls = split_url_mapping()
        self.assertIn('bsrm-steels-ltd', candidate_urls('BSRMSTEEL','BSRM Steels Limited',urls)[0])

    def test_imported_split_does_not_duplicate_manual_verified_event(self):
        import split_history
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            d = root / 'data' / 'BSRMSTEEL'
            d.mkdir(parents=True)
            (d/'company.json').write_text(json.dumps({'name':'BSRM Steels','industry':'Steel'}))
            csv_rows(d/'BSRMSTEEL_investing_splits.csv',['date','ratio','source_url'],[
                {'date':'2018-10-22', 'ratio':'1.1',
                 'source_url':'https://www.investing.com/equities/bsrm-steels-ltd-historical-data-splits'},
                {'date':'2017-10-09', 'ratio':'1.1',
                 'source_url':'https://www.investing.com/equities/bsrm-steels-ltd-historical-data-splits'}])
            verified = root/'actions.csv'
            csv_rows(verified,['ticker','event_type','date','ratio','offer_price','source_url'],[
                {'ticker':'BSRMSTEEL','event_type':'split','date':'2018-10-22',
                 'ratio':'1.1:1','offer_price':'',
                 'source_url':'https://www.investing.com/equities/bsrm-steels-ltd-historical-data-splits'}])
            with patch.object(ex,'DATA_DIR',root/'data'), patch.object(split_history,'DATA_DIR',root/'data'):
                output=ex.create_excel(2017,2018,['BSRMSTEEL'],root/'out.xlsx',actions_path=verified)
                book=load_workbook(output)
                sheet=book['Split']
                # Oct 2017: row 12; Oct 2018: row 24. Exactly 1.1, not 1.21.
                self.assertAlmostEqual(sheet['C12'].value,1.1)
                self.assertAlmostEqual(sheet['C24'].value,1.1)
                self.assertIsNone(sheet['C3'].value)


class BundledSplitFallbackTests(unittest.TestCase):
    def test_bsrm_verified_2017_2018_events_export_without_network(self):
        with tempfile.TemporaryDirectory() as folder:
            data=Path(folder)/'data'
            company=data/'BSRMSTEEL'
            company.mkdir(parents=True)
            (company/'company.json').write_text(json.dumps({'name':'BSRM Steels Limited','industry':'Steel'}))
            actions=Path(folder)/'empty_actions.csv'
            actions.write_text('ticker,event_type,date,ratio,offer_price,source_url,notes\n')
            with patch.object(ex,'DATA_DIR',data):
                out=ex.create_excel(2017,2018,['BSRMSTEEL'],Path(folder)/'bsrm.xlsx',actions_path=actions)
            sheet=load_workbook(out)['Split']
            self.assertAlmostEqual(sheet['C12'].value,1.1)
            self.assertAlmostEqual(sheet['C24'].value,1.1)
            self.assertIsNone(sheet['C3'].value)


class StockAnalysisSplitTests(unittest.TestCase):
    def test_parses_last_split_from_stockanalysis_html(self):
        from split_history import parse_stockanalysis_stats
        html = """<html><head><script>var bad='Last Split Date Jan 1, 2001 Split Ratio 200';</script></head><body>
        <h1>BSRMLTD Statistics</h1><h2>Stock Splits</h2>
        <p>The last stock split was on October 22, 2018. It was a forward split with a ratio of 1.1.</p>
        <table><tr><td>Last Split Date</td><td>Oct 22, 2018</td></tr>
        <tr><td>Split Type</td><td>Forward</td></tr>
        <tr><td>Split Ratio</td><td>1.1</td></tr></table>
        <h2>Scores</h2><p>Unrelated data</p></body></html>"""
        self.assertEqual(parse_stockanalysis_stats(html), {'date':'2018-10-22','ratio':'1.1'})
        self.assertIsNone(parse_stockanalysis_stats('<h2>Stock Splits</h2><p>No splits</p>'))
        self.assertIsNone(parse_stockanalysis_stats('<h2>Stock Splits</h2><p>Last Split Date Oct 22, 2018</p>'))

    def test_bsrm_ltd_is_not_bsrm_steel_and_history_exported(self):
        from split_history import saved_split_events
        events = saved_split_events(['BSRMLTD'])
        self.assertEqual([(e['date'].isoformat(), e['ratio']) for e in sorted(events,key=lambda e:e['date'])],
                         [('2016-03-06',1.1),('2017-09-10',1.1),('2018-10-22',1.1)])
        self.assertEqual(len(saved_split_events(['BSRMSTEEL'])),6)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root/'data'/'BSRMLTD').mkdir(parents=True)
            (root/'data'/'BSRMLTD'/'company.json').write_text(json.dumps({'name':'Bangladesh Steel Re-Rolling Mills Limited', 'industry':'Steel'}))
            actions=root/'actions.csv'
            actions.write_text('ticker,event_type,date,ratio,offer_price,source_url,notes\n')
            with patch.object(ex, 'DATA_DIR', root/'data'):
                ws=load_workbook(ex.create_excel(2016,2018,['BSRMLTD'],root/'out.xlsx',actions_path=actions))['Split']
            self.assertAlmostEqual(ws['C5'].value,1.1)   # 2016 March
            self.assertAlmostEqual(ws['C23'].value,1.1)  # 2017 September
            self.assertAlmostEqual(ws['C36'].value,1.1)  # 2018 October
            self.assertIsNone(ws['C12'].value)

    def test_refresh_gets_latest_from_stockanalysis_no_investing(self):
        import asyncio
        import split_history
        from unittest.mock import AsyncMock, MagicMock
        class Response:
            status=200
            url='https://stockanalysis.com/quote/dse/ABCTICK/statistics/'
            async def text(self):
                return '<h2>Stock Splits</h2><table><tr><td>Last Split Date</td><td>May 18, 2026</td></tr><tr><td>Split Ratio</td><td>1.15</td></tr></table>'
        request=MagicMock()
        request.get=AsyncMock(return_value=Response())
        context=MagicMock(request=request)
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(split_history,'DATA_DIR',Path(folder)):
                result=asyncio.run(split_history._fetch_stockanalysis_last_split(context,'ABCTICK',refresh=True))
                self.assertFalse(result['error'])
                request.get.assert_awaited_once_with('https://stockanalysis.com/quote/dse/ABCTICK/statistics/',timeout=22000)
                latest=split_history.saved_split_events(['ABCTICK'])
                self.assertEqual([(str(e['date']),e['ratio']) for e in latest],[('2026-05-18',1.15)])

    def test_http_403_retains_known_history_instead_of_failing(self):
        import asyncio
        import split_history
        from unittest.mock import AsyncMock, MagicMock
        blocked=MagicMock(status=403)
        request=MagicMock(get=AsyncMock(return_value=blocked))
        context=MagicMock(request=request)
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(split_history,'DATA_DIR',Path(folder)):
                result=asyncio.run(split_history._fetch_stockanalysis_last_split(context,'BSRMLTD',refresh=True))
                self.assertIn('403',result['error'])
                self.assertEqual(len(split_history.saved_split_events(['BSRMLTD'])),3)


if __name__ == '__main__':
    unittest.main()
